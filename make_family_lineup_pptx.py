# -*- coding: utf-8 -*-
"""
대가족별 반 편성 (16:9, 1장)
대가족 1~4 별로 소속 반·담당교사·학생 명단(학년)을 한 장에 압축 표기
"""
import os, sys
from datetime import date
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
sys.stdout.reconfigure(encoding='utf-8')

CELL_ORDER = ['1반','2반','3반','4반','5반','6반','고3-1','고3-2','고3-3','고3-4']

CELL_TEACHERS = {
    '1반':   {'homeroom':'이유성', 'sub':'고은솔'},
    '2반':   {'homeroom':'이윤정', 'sub':'박해니'},
    '3반':   {'homeroom':'정유빈', 'sub':'송진우'},
    '4반':   {'homeroom':'박대철', 'sub':'김희원'},
    '5반':   {'homeroom':'김광현', 'sub':'김주영'},
    '6반':   {'homeroom':'석준원', 'sub':'오은규'},
    '고3-1': {'homeroom':'이규리', 'sub':''},
    '고3-2': {'homeroom':'전성희', 'sub':''},
    '고3-3': {'homeroom':'최희승', 'sub':''},
    '고3-4': {'homeroom':'양지현', 'sub':''},
}

CELL_STUDENTS = {
    '1반': [('오은수',1,'여'),('김나현',2,'여'),('김채은',2,'여'),('박지웅',1,'남'),
            ('이가영',2,'여'),('민예원',2,'여'),('이서현',1,'여'),('이소현',1,'여'),('김라일',2,'남')],
    '2반': [('김정록',1,'남'),('안민혁',1,'남'),('박민주',1,'여'),('전지훈',2,'남'),
            ('이주형',1,'여'),('박동하',1,'남'),('김강은',1,'여'),('원대연',2,'남'),('유재희',2,'여')],
    '3반': [('김한결',1,'남'),('최지은',1,'여'),('최민호',2,'남'),('김강희',1,'여'),
            ('배유나',2,'여'),('박재영',1,'남'),('조한나',1,'여'),('하은유',1,'여'),('장지현',2,'여')],
    '4반': [('고예나',1,'여'),('이빛가온',1,'남'),('최수지',1,'여'),('윤서희',2,'여'),
            ('김재원',1,'남'),('하어림',1,'남'),('오현우',1,'남'),('연해준',2,'남'),('송준원',1,'남')],
    '5반': [('윤성하',1,'남'),('김은',1,'여'),('김유진',2,'여'),('임태민',2,'남'),
            ('강은교',2,'여'),('유나린',2,'여'),('권하율',1,'남'),('김주하',2,'여'),('정태율',1,'남')],
    '6반': [('나호윤',1,'남'),('이서후',1,'여'),('최유현',1,'여'),('김도연',2,'여'),
            ('양소희',2,'여'),('홍기진',2,'남'),('이준호',2,'남'),('박시우',1,'남'),('이하준',1,'남')],
    '고3-1': [('송성모',3,'남'),('김도헌',3,'남'),('이소민',3,'여'),('박지호',3,'여'),
              ('한채은',3,'여'),('김준수',3,'남'),('박진서',3,'여'),('임소명',3,'여')],
    '고3-2': [('이에녹',3,'남'),('백결',3,'남'),('조혁준',3,'남'),('최승아',3,'여'),
              ('조강인',3,'남'),('정승빈',3,'여'),('서이수',3,'여'),('김민채',3,'여')],
    '고3-3': [('곽지원',3,'남'),('김다은',3,'여'),('장세민',3,'남'),('정원영',3,'남'),
              ('주성',3,'남'),('나정인',3,'여'),('명주원',3,'남')],
    '고3-4': [('최가은',3,'여'),('김은우',3,'여'),('최지윤',3,'여'),('최서윤',3,'여'),
              ('오희수',3,'여'),('심희은',3,'여'),('조인상',3,'여'),('양인혁',3,'남')],
}

CELL_FAMILY = {'1반':1,'2반':1,'3반':2,'4반':2,'5반':3,'6반':3,
               '고3-1':4,'고3-2':4,'고3-3':4,'고3-4':4}
FAMILY_LEADERS = {1:'이윤정 (2반)', 2:'박대철 (4반)', 3:'석준원 (6반)', 4:'전성희 (고3-2)'}

DARK   = RGBColor(0x17,0x13,0x2B)
LILAC  = RGBColor(0xA7,0x8B,0xFA)
GREEN  = RGBColor(0x16,0xA3,0x74)
ORANGE = RGBColor(0xEA,0x58,0x0C)
FEMALE = RGBColor(0xBE,0x18,0x5D)
INK    = RGBColor(0x1F,0x29,0x37)
MUTED  = RGBColor(0x6B,0x72,0x80)
WHITE  = RGBColor(0xFF,0xFF,0xFF)
MALE_C = RGBColor(0x1D,0x4E,0xD8)
FEMALE_C = FEMALE
FAM_LINE = {1:LILAC, 2:GREEN, 3:ORANGE, 4:FEMALE}
FAM_BG   = {1:RGBColor(0xF5,0xF3,0xFF), 2:RGBColor(0xEC,0xFD,0xF5),
            3:RGBColor(0xFF,0xF7,0xED), 4:RGBColor(0xFE,0xF2,0xF7)}
FONT = '맑은 고딕'

def sort_st(students):
    return sorted(students, key=lambda t: (-t[1], t[0]))

def runs(p, parts):
    for txt, size, bold, color in parts:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = FONT

def family_box(slide, x, y, w, h, fam, classes):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = FAM_BG[fam]
    box.line.color.rgb = FAM_LINE[fam]; box.line.width = Pt(1.5)
    box.adjustments[0] = 0.045; box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.08); tf.margin_bottom = Inches(0.05)

    total_n = sum(len(CELL_STUDENTS[c]) for c in classes)
    p = tf.paragraphs[0]
    runs(p, [(f'대가족 {fam}', 16, True, FAM_LINE[fam]), (f'   총 {total_n}명', 10, False, MUTED),
             (f'   ★가족장 {FAMILY_LEADERS[fam]}', 10, True, INK)])

    for cell in classes:
        t = CELL_TEACHERS[cell]
        is_g3 = cell.startswith('고3')
        students = sort_st(CELL_STUDENTS[cell])
        p1 = tf.add_paragraph(); p1.space_before = Pt(5)
        teacher_txt = f'담임 {t["homeroom"]}' if is_g3 else f'교사1 {t["homeroom"]} · 교사2 {t["sub"]}'
        runs(p1, [(f'{cell}', 12.5, True, INK), (f'  {teacher_txt}', 9.5, False, MUTED),
                   (f'  ({len(students)}명)', 9, False, MUTED)])
        p2 = tf.add_paragraph(); p2.space_before = Pt(1)
        name_parts = []
        for i, (name, grade, gender) in enumerate(students):
            sep = ' · ' if i > 0 else ''
            name_parts.append((sep + name, 9.5, False, INK))
            name_parts.append((f'({grade})', 8, False, MALE_C if gender == '남' else FEMALE_C))
        runs(p2, name_parts)

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE; bg.line.fill.background(); bg.shadow.inherit = False

    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(9), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    runs(p, [('👨‍👩‍👧‍👦 대가족별 반 편성', 24, True, DARK)])

    rb = s.shapes.add_textbox(Inches(9.8), Inches(0.26), Inches(3.1), Inches(0.6))
    rtf = rb.text_frame
    rr = rtf.paragraphs[0]; rr.alignment = PP_ALIGN.RIGHT
    rn = rr.add_run(); rn.text = f'내수동 경주자 2026 · {date.today().strftime("%Y-%m-%d")}'
    rn.font.size = Pt(11); rn.font.bold = True; rn.font.color.rgb = LILAC; rn.font.name = FONT

    top = Inches(1.0)
    bottom_margin = Inches(0.35)
    left_x = Inches(0.5)
    box_w = prs.slide_width - left_x * 2
    total_h = prs.slide_height - top - bottom_margin
    gap = Inches(0.18)

    fam_classes = {1: [], 2: [], 3: [], 4: []}
    for c in CELL_ORDER:
        fam_classes[CELL_FAMILY[c]].append(c)
    counts = {f: len(cs) for f, cs in fam_classes.items()}
    unit_h = (total_h - gap * 3) / sum(counts.values())

    y = top
    for fam in [1, 2, 3, 4]:
        h = unit_h * counts[fam]
        family_box(s, left_x, y, box_w, h, fam, fam_classes[fam])
        y += h + gap

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'family-lineup.pptx')
    prs.save(out)
    print('family-lineup.pptx 생성 완료:', out, f'({os.path.getsize(out)/1024:.0f} KB)')

if __name__ == '__main__':
    main()
