# -*- coding: utf-8 -*-
"""
고등부 전체 일정 유인물 (16:9, 1장)
Supabase schedule 테이블 기준 (2026-09 ~ 2027-01, 22건)
"""
import os, sys
from datetime import date
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
sys.stdout.reconfigure(encoding='utf-8')

# ===== 데이터 (Supabase schedule 테이블과 동일 — 학기 중 갱신 시 이 배열도 갱신) =====
SCHEDULE = [
    ('2026-09-06', '소그룹', '교사 전체회의 · 신규텀 시작'),
    ('2026-09-13', '소그룹', '국장단 회의'),
    ('2026-09-20', '소그룹 아웃팅', None),
    ('2026-09-27', '대가족반 모임', '9/24(목)~9/26(토) 추석 연휴 · 결석자 많을 것으로 예상'),
    ('2026-10-04', '소그룹', '교사 전체회의 · 10/5(월) 대체공휴일'),
    ('2026-10-11', '소그룹', '국장단 회의 · 중간고사 시즌 가능성 높음'),
    ('2026-10-18', '소그룹', None),
    ('2026-10-25', '학년모임', None),
    ('2026-11-01', '소그룹', '교사 전체회의 · 수능 전 주간'),
    ('2026-11-08', '소그룹', '국장단 회의 · 수능 전 주간'),
    ('2026-11-15', '정기예배 + 고3 기도회', '11/19(목) 수능'),
    ('2026-11-22', '소그룹', None),
    ('2026-11-29', '소그룹 아웃팅', None),
    ('2026-12-06', '소그룹', '교사 전체회의 · 고3 졸업여행 후보 A'),
    ('2026-12-13', '소그룹', '국장단 회의 · 고3 졸업여행 후보 B'),
    ('2026-12-20', '소그룹', '12/25(금) 성탄절 (학년모임 운영 여부 고민 중)'),
    ('2026-12-27', '졸업예배 + 소그룹 아웃팅', None),
    ('2027-01-03', '소그룹', '교사 전체회의 · 고1 유입 · 신규텀 시작'),
    ('2027-01-10', '학부모 간담회 / 학년모임', '국장단 회의 · 학년별 수련회 특송 준비'),
    ('2027-01-17', '겨울수련회 후보 A', '1/15(금)~1/17(일)'),
    ('2027-01-24', '겨울수련회 후보 B', '1/22(금)~1/24(일)'),
    ('2027-01-31', '소그룹 아웃팅', None),
]
MONTH_LABEL = {9: '9월', 10: '10월', 11: '11월', 12: '12월', 1: '1월'}
COLUMNS = [[9, 10], [11, 12], [1]]  # 3열 분배

DARK   = RGBColor(0x17,0x13,0x2B)
LILAC  = RGBColor(0xA7,0x8B,0xFA)
PURPLE = RGBColor(0x7C,0x3A,0xED)
INK    = RGBColor(0x1F,0x29,0x37)
MUTED  = RGBColor(0x6B,0x72,0x80)
WHITE  = RGBColor(0xFF,0xFF,0xFF)
MEET_COLOR = RGBColor(0xB4,0x53,0x09)
FONT = '맑은 고딕'

def runs(p, parts):
    for txt, size, bold, color, italic in parts:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = FONT

def month_of(d): return int(d.split('-')[1])
def day_label(d): return str(int(d.split('-')[1])) + '/' + str(int(d.split('-')[2]))

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE; bg.line.fill.background(); bg.shadow.inherit = False

    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(8.5), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    runs(p, [('📅 고등부 전체 일정', 26, True, DARK, False)])
    p2 = tb.text_frame.add_paragraph(); p2.space_before = Pt(2)
    runs(p2, [('2026년 9월 ~ 2027년 1월 (4학기)', 13, False, MUTED, False)])

    rb = s.shapes.add_textbox(Inches(9.6), Inches(0.32), Inches(3.3), Inches(0.7))
    rtf = rb.text_frame
    rr = rtf.paragraphs[0]; rr.alignment = PP_ALIGN.RIGHT
    rn = rr.add_run(); rn.text = '내수동 경주자 2026'
    rn.font.size = Pt(12); rn.font.bold = True; rn.font.color.rgb = LILAC; rn.font.name = FONT
    rr2 = rtf.add_paragraph(); rr2.alignment = PP_ALIGN.RIGHT
    rn2 = rr2.add_run(); rn2.text = date.today().strftime('%Y-%m-%d 발행')
    rn2.font.size = Pt(11); rn2.font.color.rgb = MUTED; rn2.font.name = FONT

    top = Inches(1.35)
    bottom_margin = Inches(0.55)
    left = Inches(0.5)
    total_w = prs.slide_width - left * 2
    gap = Inches(0.35)
    col_w = (total_w - gap * 2) / 3
    col_h = prs.slide_height - top - bottom_margin

    for i, months in enumerate(COLUMNS):
        x = left + i * (col_w + gap)
        box = s.shapes.add_textbox(x, top, col_w, col_h)
        tf = box.text_frame; tf.word_wrap = True
        first_month = True
        for m in months:
            rows = [r for r in SCHEDULE if month_of(r[0]) == m]
            if not rows:
                continue
            if first_month:
                mp = tf.paragraphs[0]; first_month = False
            else:
                mp = tf.add_paragraph(); mp.space_before = Pt(12)
            runs(mp, [(MONTH_LABEL[m], 16, True, PURPLE, False)])
            for d, title, note in rows:
                rp = tf.add_paragraph(); rp.space_before = Pt(6)
                runs(rp, [(day_label(d) + '  ', 12.5, True, INK, False), (title, 12.5, True, INK, False)])
                if note:
                    np_ = tf.add_paragraph(); np_.space_before = Pt(0)
                    is_meeting = ('전체회의' in note) or ('국장단 회의' in note)
                    color = MEET_COLOR if is_meeting else MUTED
                    runs(np_, [('     ' + note, 10.5, False, color, not is_meeting)])

    # 범례
    leg = s.shapes.add_textbox(Inches(0.5), prs.slide_height - Inches(0.45), Inches(10), Inches(0.35))
    lp = leg.text_frame.paragraphs[0]
    runs(lp, [('매월 첫째 주일 = 교사 전체회의 · 매월 둘째 주일 = 국장단 회의', 10.5, False, MEET_COLOR, False)])

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'schedule-handout.pptx')
    prs.save(out)
    print('schedule-handout.pptx 생성 완료:', out)

if __name__ == '__main__':
    main()
