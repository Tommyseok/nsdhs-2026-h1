# -*- coding: utf-8 -*-
"""
공과공부 반별 배치도 (16:9, 1장)
장소 4곳: 강당(6개반) · 유리방(2개반) · 교사방(1개반) · 준2층(1개반)
대가족 단위로 묶어서 배치 (대가족1~3 = 강당, 대가족4 = 유리방+교사방+준2층)
"""
import os, sys
from datetime import date
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
sys.stdout.reconfigure(encoding='utf-8')

CELL_TEACHERS = {
    '1반':   {'homeroom':'이유성', 'sub':'고은솔'},
    '2반':   {'homeroom':'이윤정', 'sub':'박해니'},
    '3반':   {'homeroom':'정유빈', 'sub':'송진우'},
    '4반':   {'homeroom':'박대철', 'sub':'김희원'},
    '5반':   {'homeroom':'김광현', 'sub':'김주영'},
    '6반':   {'homeroom':'석준원', 'sub':'오은규'},
    '고3-1': {'homeroom':'이규리', 'sub':''},
    '고3-2': {'homeroom':'전성희', 'sub':''},
    '고3-3': {'homeroom':'최희승', 'sub':''},
    '고3-4': {'homeroom':'양지현', 'sub':''},
}
CELL_FAMILY = {'1반':1,'2반':1,'3반':2,'4반':2,'5반':3,'6반':3,
               '고3-1':4,'고3-2':4,'고3-3':4,'고3-4':4}

# 방 배치: 대가족 단위로 묶임 (강당=대가족1~3 전체, 유리방+교사방+준2층=대가족4 전체)
ROOMS = [
    {'name':'강당',   'sub':'대가족 1·2·3', 'classes':['1반','2반','3반','4반','5반','6반'], 'cols':2},
    {'name':'유리방', 'sub':'대가족 4',      'classes':['고3-1','고3-2'], 'cols':1},
    {'name':'교사방', 'sub':'대가족 4',      'classes':['고3-3'], 'cols':1},
    {'name':'준2층',  'sub':'대가족 4',      'classes':['고3-4'], 'cols':1},
]

DARK   = RGBColor(0x17,0x13,0x2B)
LILAC  = RGBColor(0xA7,0x8B,0xFA)
GREEN  = RGBColor(0x16,0xA3,0x74)
ORANGE = RGBColor(0xEA,0x58,0x0C)
FEMALE = RGBColor(0xBE,0x18,0x5D)
INK    = RGBColor(0x1F,0x29,0x37)
MUTED  = RGBColor(0x6B,0x72,0x80)
WHITE  = RGBColor(0xFF,0xFF,0xFF)
ROOM_BORDER = RGBColor(0xCB,0xD5,0xE1)
ROOM_HEAD_BG= RGBColor(0xF1,0xF5,0xF9)
FAM_LINE = {1:LILAC, 2:GREEN, 3:ORANGE, 4:FEMALE}
FAM_BG   = {1:RGBColor(0xF5,0xF3,0xFF), 2:RGBColor(0xEC,0xFD,0xF5),
            3:RGBColor(0xFF,0xF7,0xED), 4:RGBColor(0xFE,0xF2,0xF7)}
FONT = '맑은 고딕'

def runs(p, parts):
    for txt, size, bold, color in parts:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = FONT

def class_chip(slide, x, y, w, h, cell):
    fam = CELL_FAMILY[cell]
    t = CELL_TEACHERS[cell]
    is_g3 = cell.startswith('고3')
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = FAM_BG[fam]
    box.line.color.rgb = FAM_LINE[fam]; box.line.width = Pt(1.25)
    box.adjustments[0] = 0.10; box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]
    runs(p, [(cell, 17, True, INK), (f'  대가족{fam}', 10, True, FAM_LINE[fam])])
    p2 = tf.add_paragraph(); p2.space_before = Pt(2)
    if is_g3:
        runs(p2, [('담임 ', 11, False, MUTED), (t['homeroom'], 12, True, INK)])
    else:
        runs(p2, [('교사1 ', 11, False, MUTED), (t['homeroom'], 12, True, INK),
                   ('   교사2 ', 11, False, MUTED), (t['sub'], 12, True, INK)])

def room_box(slide, x, y, w, h, room):
    outer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    outer.fill.solid(); outer.fill.fore_color.rgb = WHITE
    outer.line.color.rgb = ROOM_BORDER; outer.line.width = Pt(1.5)
    outer.adjustments[0] = 0.04; outer.shadow.inherit = False
    outer.text_frame.paragraphs[0].add_run().text = ''

    head_h = Inches(0.55)
    head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, head_h)
    head.fill.solid(); head.fill.fore_color.rgb = ROOM_HEAD_BG
    head.line.fill.background(); head.adjustments[0] = 0.35; head.shadow.inherit = False
    htf = head.text_frame; htf.vertical_anchor = MSO_ANCHOR.MIDDLE
    htf.margin_left = Inches(0.2)
    hp = htf.paragraphs[0]
    runs(hp, [(f'📍 {room["name"]}', 17, True, DARK), (f'   ({room["sub"]} · {len(room["classes"])}개반)', 12, False, MUTED)])

    cols = room['cols']
    n = len(room['classes'])
    rows = (n + cols - 1) // cols
    pad = Inches(0.18)
    gap = Inches(0.15)
    area_x = x + pad
    area_y = y + head_h + pad
    area_w = w - pad * 2
    area_h = h - head_h - pad * 2
    cw = (area_w - gap * (cols - 1)) / cols
    ch = (area_h - gap * (rows - 1)) / rows
    for i, cell in enumerate(room['classes']):
        r_i, c_i = divmod(i, cols)
        cx = area_x + c_i * (cw + gap)
        cy = area_y + r_i * (ch + gap)
        class_chip(slide, cx, cy, cw, ch, cell)

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE; bg.line.fill.background(); bg.shadow.inherit = False

    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(9), Inches(0.8))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    runs(p, [('📖 공과공부 반별 배치도', 26, True, DARK)])

    rb = s.shapes.add_textbox(Inches(9.8), Inches(0.32), Inches(3.1), Inches(0.7))
    rtf = rb.text_frame
    rr = rtf.paragraphs[0]; rr.alignment = PP_ALIGN.RIGHT
    rn = rr.add_run(); rn.text = '내수동 경주자 2026'
    rn.font.size = Pt(12); rn.font.bold = True; rn.font.color.rgb = LILAC; rn.font.name = FONT
    rr2 = rtf.add_paragraph(); rr2.alignment = PP_ALIGN.RIGHT
    rn2 = rr2.add_run(); rn2.text = date.today().strftime('%Y-%m-%d')
    rn2.font.size = Pt(11); rn2.font.color.rgb = MUTED; rn2.font.name = FONT

    # 좌측 큰 방: 강당 (2열 x 3행) / 우측: 유리방·교사방·준2층 세로 스택
    top = Inches(1.25)
    bottom_margin = Inches(0.4)
    total_h = prs.slide_height - top - bottom_margin

    left_x = Inches(0.5)
    left_w = Inches(7.6)
    room_box(s, left_x, top, left_w, total_h, ROOMS[0])

    right_x = left_x + left_w + Inches(0.3)
    right_w = prs.slide_width - right_x - Inches(0.5)
    gap = Inches(0.25)
    # 방마다 헤더 고정 오버헤드 + 반 개수에 비례한 높이를 배정 (칩 텍스트 겹침 방지)
    HEADER_UNIT = 0.5
    CLASS_UNIT = 1.3
    weights = [HEADER_UNIT + len(r['classes']) * CLASS_UNIT for r in ROOMS[1:]]
    total_weight = sum(weights)
    avail_h = total_h - gap * (len(ROOMS[1:]) - 1)
    y = top
    for room, wgt in zip(ROOMS[1:], weights):
        r_h = avail_h * (wgt / total_weight)
        room_box(s, right_x, y, right_w, r_h, room)
        y += r_h + gap

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'room-layout.pptx')
    prs.save(out)
    print('room-layout.pptx 생성 완료:', out, f'({os.path.getsize(out)/1024:.0f} KB)')

if __name__ == '__main__':
    main()
